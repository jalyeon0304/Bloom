from collections import deque
from copy import deepcopy
from datetime import UTC, datetime, timedelta
import json
import re
import subprocess
from pathlib import Path
from secrets import token_urlsafe
from uuid import uuid4
from zoneinfo import ZoneInfo

from fastapi import FastAPI, Header, HTTPException, Query
from fastapi.responses import HTMLResponse, Response
from pydantic import BaseModel, Field

from bloom.config import settings
from bloom.demo_data import (
    activate_site,
    list_master_sites,
    list_sites,
    series_for_site,
    session_status,
    start_session,
)

app = FastAPI(title=settings.app_name)

UI_MARKER_SUMMARY = "summary-main-v1"
UI_MARKER_SITE_DETAIL = "site-detail-v1"


def _git_cmd(args: list[str]) -> str | None:
    try:
        return subprocess.check_output(args, cwd=Path(__file__).parents[2], text=True).strip()
    except Exception:
        return None


def _build_info() -> dict[str, str | bool]:
    sha = _git_cmd(["git", "rev-parse", "--short", "HEAD"]) or "unknown"
    branch = _git_cmd(["git", "rev-parse", "--abbrev-ref", "HEAD"]) or "unknown"
    status = _git_cmd(["git", "status", "--porcelain"])
    dirty = bool(status) if status is not None else False
    return {"sha": sha, "branch": branch, "dirty": dirty}


class LoginRequestIn(BaseModel):
    username: str = Field(min_length=2, max_length=50)
    password: str = Field(min_length=4, max_length=128)
    device_name: str = Field(default="unknown-device", max_length=80)


class ApproveIn(BaseModel):
    approver: str = Field(default="admin", max_length=50)


class TokenExchangeIn(BaseModel):
    request_id: str


class SiteActivationIn(BaseModel):
    site_id: str = Field(min_length=2, max_length=20)
    submitted_output_kw: int = Field(gt=0)
    min_power_kw: int | None = Field(default=None, ge=0)
    max_power_kw: int | None = Field(default=None, ge=0)
    target_pct: int | None = Field(default=None, ge=0, le=100)
    deadline_minutes: int | None = Field(default=None, gt=0)
    semi_target_enabled: bool = False
    min_target_minutes: int | None = Field(default=None, gt=0)
    max_target_minutes: int | None = Field(default=None, gt=0)


class SchedulePatchIn(BaseModel):
    date: str | None = None
    siteId: str
    field: str
    value: str | int | bool


class UiKitSchedulePatchIn(BaseModel):
    siteId: str
    field: str
    value: str | int | bool


# demo purpose only: in-memory store (replace with DB/Redis in production)
_pending_login_requests: dict[str, dict] = {}
_session_tokens: dict[str, dict] = {}

_SCHEDULE_ROWS: list[dict[str, str]] = [
    {"site_id": "SKK167", "field": "target_kw", "value": "12000", "status": "active"},
    {"site_id": "SKK046", "field": "min_power_kw", "value": "10200", "status": "active"},
    {"site_id": "KMP000", "field": "target_kw", "value": "2400", "status": "cancel"},
]
_SCHEDULE_LOGS: list[str] = []

_UIKIT_SCENARIO_FILES = {
    "normal": "summary_normal.json",
    "loading": "summary_loading.json",
    "empty": "summary_empty.json",
    "activeHighCapacity": "summary_activeHighCapacity.json",
}


def _load_uikit_state() -> dict[str, dict]:
    base = Path(__file__).with_name("ui_fixtures")
    out: dict[str, dict] = {}
    for scenario, filename in _UIKIT_SCENARIO_FILES.items():
        fixture_path = base / filename
        out[scenario] = json.loads(fixture_path.read_text(encoding="utf-8"))
        out[scenario].setdefault("systemLogs", [])
    return out


_UIKIT_DATA: dict[str, dict] = _load_uikit_state()


def _filter_skk_payload(data: dict) -> dict:
    pattern = re.compile(r"^SKK\d{3}$")
    payload = deepcopy(data)
    payload["sites"] = [s for s in payload.get("sites", []) if pattern.match(str(s.get("siteId") or ""))]
    tmap = payload.get("timeseriesBySiteId") or {}
    payload["timeseriesBySiteId"] = {k: v for k, v in tmap.items() if pattern.match(str(k))}
    payload["series"] = [p for p in payload.get("series", [])]
    if payload.get("semiTable"):
        payload["semiTable"] = [r for r in payload.get("semiTable", []) if pattern.match(str(r.get("siteId") or ""))]
    if payload.get("nonTable"):
        payload["nonTable"] = [r for r in payload.get("nonTable", []) if pattern.match(str(r.get("siteId") or ""))]
    return payload


def _make_uikit_logline(site_id: str, field: str, old: str, new: str, user: str | None) -> str:
    stamp = _now_kst().strftime("%H:%M:%S")
    if user:
        return f"[{stamp}] {site_id} {field} : {old} → {new} (user:{user})"
    return f"[{stamp}] {site_id} {field} : {old} → {new}"



def _role_guard(x_role: str | None, allowed: set[str]) -> str:
    role = (x_role or "viewer").strip().lower()
    if role not in allowed:
        raise HTTPException(status_code=403, detail="forbidden")
    return role


def _now_kst() -> datetime:
    return datetime.now(tz=ZoneInfo("Asia/Seoul"))


def _normalize_log_value(field: str, value: str | int | bool | None) -> str:
    if value is None:
        return ""
    if field in {"rmccStart", "start", "finish"}:
        text = str(value).strip()
        m = re.match(r"^(\d{1,2}):(\d{1,2})$", text)
        if m:
            return f"{int(m.group(1)):02d}:{int(m.group(2)):02d}"
        return text
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def _order_label(site_map: dict) -> str:
    return str(site_map.get("customerOrder") or site_map.get("orderLabel") or "KPX Rampdown")


def _make_operation_log(
    *,
    date_text: str,
    site_id: str,
    site_map: dict,
    field: str,
    old_value: str | int | bool | None,
    new_value: str | int | bool,
) -> dict:
    timestamp_kst = f"{date_text} {_now_kst().strftime('%H:%M')}"
    order_label = _order_label(site_map)
    if field == "isActive":
        event_type = "ACTIVATE" if bool(new_value) else "DEACTIVATE"
        detail = f"{order_label} {'Activate' if bool(new_value) else 'Deactivate'}"
    elif field == "isCanceled":
        event_type = "CANCELED" if bool(new_value) else "CANCEL_CLEARED"
        detail = f"{order_label} {'Canceled' if bool(new_value) else 'Cancel cleared'}"
    else:
        event_type = "CHANGED"
        old_norm = _normalize_log_value(field, old_value) or "∅"
        new_norm = _normalize_log_value(field, new_value) or "∅"
        detail = f"{order_label} changed ({field}: {old_norm} → {new_norm})"
    message = f"[{timestamp_kst}] {site_id} {detail}"
    return {
        "timestampKst": timestamp_kst,
        "siteId": site_id,
        "eventType": event_type,
        "detail": detail,
        "message": message,
    }


def _normalize_operation_logs(logs: list) -> list[dict]:
    normalized: list[dict] = []
    for item in logs:
        if isinstance(item, dict) and item.get("message"):
            normalized.append(
                {
                    "timestampKst": str(item.get("timestampKst") or ""),
                    "siteId": str(item.get("siteId") or ""),
                    "eventType": str(item.get("eventType") or "CHANGED"),
                    "detail": str(item.get("detail") or ""),
                    "message": str(item.get("message") or ""),
                }
            )
        elif isinstance(item, str):
            normalized.append(
                {
                    "timestampKst": "",
                    "siteId": "",
                    "eventType": "CHANGED",
                    "detail": item,
                    "message": item,
                }
            )
    return normalized


def _make_logline(site_id: str, field: str, old: str, new: str, user: str | None) -> str:
    stamp = _now_kst().strftime("%H:%M:%S")
    who = user or "unknown"
    return f"[{stamp}] {site_id} {field} : {old} → {new} (user:{who})"


def _has_route(path: str) -> bool:
    return any(getattr(route, "path", "") == path for route in app.routes)


@app.get("/health")
def health() -> dict:
    build = _build_info()
    return {
        "status": "ok",
        "env": settings.app_env,
        "mainModulePath": __file__,
        "dashboardUiMarker": UI_MARKER_SUMMARY,
        "siteDetailUiMarker": UI_MARKER_SITE_DETAIL,
        "hasSummaryRoute": _has_route("/summary"),
        "hasSiteDetailRoute": _has_route("/sites/{site_id}"),
        "gitSha": build["sha"],
        "gitBranch": build["branch"],
        "dirty": build["dirty"],
    }


@app.get("/api/master-sites")
def api_master_sites(group: str | None = Query(default=None)) -> dict:
    return {"items": list_master_sites(group)}


@app.post("/api/sites/activate")
def api_activate_site(payload: SiteActivationIn) -> dict:
    try:
        return activate_site(
            payload.site_id,
            payload.submitted_output_kw,
            payload.min_power_kw,
            payload.max_power_kw,
            payload.target_pct,
            payload.deadline_minutes,
            payload.semi_target_enabled,
            payload.min_target_minutes,
            payload.max_target_minutes,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/session/start")
def api_start_session() -> dict:
    return start_session()


@app.get("/api/session/status")
def api_session_status() -> dict:
    return session_status()


@app.get("/api/sites")
def api_sites(group: str | None = Query(default=None)) -> dict:
    return {"items": list_sites(group)}


@app.get("/api/series/{site_id}")
def api_series(site_id: str, points: int = Query(default=30, ge=10, le=180)) -> dict:
    series = series_for_site(site_id, points)
    if not series:
        raise HTTPException(status_code=404, detail="site not found")
    return {"site_id": site_id, "points": series}


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard() -> str:
    html_path = Path(__file__).with_name("ui_summary_main.html")
    build = _build_info()
    html = html_path.read_text(encoding="utf-8").replace("__BUILD_SHA__", str(build["sha"]))
    return HTMLResponse(
        content=html,
        headers={"X-Bloom-UI": UI_MARKER_SUMMARY, "X-Bloom-Build": str(build["sha"])},
    )


@app.get("/ui-kit", response_class=HTMLResponse)
def ui_kit(scenario: str = Query(default="normal")) -> str:
    if settings.app_env not in {"dev", "local"} and not settings.bloom_debug_ui:
        raise HTTPException(status_code=404, detail="not found")
    html_path = Path(__file__).with_name("ui_summary_preview.html")
    html = html_path.read_text(encoding="utf-8")
    return html.replace("__SCENARIO__", scenario)


@app.get("/api/ui-kit/summary")
def api_ui_kit_summary(scenario: str = Query(default="normal")) -> dict:
    key = scenario if scenario in _UIKIT_DATA else "normal"
    return _filter_skk_payload(_UIKIT_DATA[key])


@app.patch("/api/ui-kit/schedule")
def api_ui_kit_schedule_patch(
    payload: UiKitSchedulePatchIn,
    scenario: str = Query(default="normal"),
    x_user_role: str | None = Header(default=None),
    x_user_name: str | None = Header(default=None),
) -> dict:
    _role_guard(x_user_role, {"operator"})

    scenario_key = scenario if scenario in _UIKIT_DATA else "normal"
    data = _filter_skk_payload(_UIKIT_DATA[scenario_key])
    _UIKIT_DATA[scenario_key] = data
    site_id = payload.siteId
    field = payload.field

    semi = data.get("semiTable") or []
    non = data.get("nonTable") or []
    row = next((r for r in semi + non if r.get("siteId") == site_id), None)
    if row is None:
        raise HTTPException(status_code=404, detail="site row not found")

    old = row.get(field)
    row[field] = payload.value

    for s in data.get("sites", []):
        if s.get("siteId") == site_id and field == "isCanceled":
            s["isCanceled"] = bool(payload.value)

    log_line = None
    if str(old) != str(payload.value):
        log_line = _make_uikit_logline(site_id, field, str(old), str(payload.value), x_user_name)
        logs = data.setdefault("systemLogs", [])
        logs.append(log_line)

    return {
        "ok": True,
        "updated": {"siteId": site_id, "field": field, "value": payload.value},
        "logLine": log_line,
    }




def _render_site_detail(site_id: str) -> HTMLResponse:
    html_path = Path(__file__).with_name("ui_site_detail.html")
    html = html_path.read_text(encoding="utf-8")
    return HTMLResponse(
        content=html.replace("__SITE_ID__", site_id),
        headers={"X-Bloom-UI": UI_MARKER_SITE_DETAIL, "X-Bloom-Site": site_id},
    )


@app.get("/sites/{site_id}", response_class=HTMLResponse)
@app.get("/sites/{site_id}/", response_class=HTMLResponse)
def site_detail(site_id: str) -> HTMLResponse:
    return _render_site_detail(site_id)


@app.get("/site/{site_id}", response_class=HTMLResponse)
def site_detail_legacy(site_id: str) -> HTMLResponse:
    return _render_site_detail(site_id)


@app.get("/summary", response_class=HTMLResponse)
def summary_page() -> str:
    return dashboard()


@app.get("/api/version")
def api_version() -> dict:
    build = _build_info()
    return {
        "dashboardUiMarker": UI_MARKER_SUMMARY,
        "siteDetailUiMarker": UI_MARKER_SITE_DETAIL,
        "hasSummaryRoute": True,
        "hasSiteDetailRoute": True,
        "mainModulePath": __file__,
        "gitSha": build["sha"],
        "gitBranch": build["branch"],
        "dirty": build["dirty"],
    }


@app.get("/api/routes")
def api_routes() -> dict:
    paths = sorted(
        {
            route.path
            for route in app.routes
            if getattr(route, "path", "").startswith("/")
        }
    )
    return {
        "count": len(paths),
        "paths": paths,
    }


@app.get("/api/meta")
def api_meta() -> dict:
    now = _now_kst()
    return {
        "timezone": "Asia/Seoul",
        "serverDateKst": now.strftime("%Y-%m-%d"),
        "baselineTime": "10:00",
        "endTime": "17:00",
    }


def _store_path() -> Path:
    return Path(__file__).with_name("schedule_store.json")


def _load_schedule_store() -> dict:
    path = _store_path()
    if not path.exists():
        return {"byDate": {}, "systemLogs": {}}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {"byDate": {}, "systemLogs": {}}
    data.setdefault("byDate", {})
    data.setdefault("systemLogs", {})
    return data


def _save_schedule_store(store: dict) -> None:
    _store_path().write_text(json.dumps(store, ensure_ascii=False, indent=2), encoding="utf-8")


def _history_path() -> Path:
    return Path(__file__).parents[2] / "data" / "dispatch_history.jsonl"


def _ensure_history_file() -> Path:
    path = _history_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.touch()
    return path


def _iter_history_records():
    path = _history_path()
    if not path.exists():
        return
    with path.open(encoding="utf-8") as fp:
        for line in fp:
            text = line.strip()
            if not text:
                continue
            try:
                yield json.loads(text)
            except json.JSONDecodeError:
                continue


def _append_history(record: dict) -> None:
    path = _ensure_history_file()
    with path.open("a", encoding="utf-8") as fp:
        fp.write(json.dumps(record, ensure_ascii=False) + "\n")


def _event_type(field: str, old_value: str | int | bool | None, new_value: str | int | bool) -> str:
    if field == "isActive" and bool(old_value) is False and new_value is True:
        return "ORDER_CREATED"
    if field == "isCanceled" and bool(old_value) is False and new_value is True:
        return "ORDER_CANCELED"
    return "ORDER_UPDATED"


def _make_snapshot_record(
    date: str,
    site: dict,
    field: str,
    old_value: str | int | bool | None,
    new_value: str | int | bool,
    role: str,
    user_name: str | None,
) -> dict:
    return {
        "eventType": _event_type(field, old_value, new_value),
        "occurredAt": _now_kst().isoformat(timespec="seconds"),
        "userRole": role,
        "userName": user_name or None,
        "date": date,
        "siteId": site.get("siteId"),
        "siteName": site.get("siteName"),
        "group": site.get("group"),
        "capacityKw": site.get("capacityKw"),
        "rmccStart": site.get("rmccStart"),
        "start": site.get("start"),
        "finish": site.get("finish"),
        "minKw": site.get("minKw"),
        "maxKw": site.get("maxKw"),
        "baselineKw1000": site.get("baselineKw1000") if site.get("baselineKw1000") is not None else site.get("baselineKw0900"),
        "baselineKw0900": site.get("baselineKw1000") if site.get("baselineKw1000") is not None else site.get("baselineKw0900"),
        "isActive": site.get("isActive"),
        "isCanceled": site.get("isCanceled"),
    }


def _server_date(date: str | None) -> str:
    return date or _now_kst().strftime("%Y-%m-%d")


def _excel_kst_datetime(date_text: str, hhmm: str | None) -> str | None:
    if not hhmm or not _valid_hhmm(hhmm):
        return None
    return f"{date_text} {hhmm}"


def _valid_hhmm(value: str) -> bool:
    if len(value) != 5 or value[2] != ":":
        return False
    hh = value[:2]
    mm = value[3:]
    if not (hh.isdigit() and mm.isdigit()):
        return False
    mins = int(hh) * 60 + int(mm)
    return 600 <= mins <= 1020


def _baseline_from_points(points: list[dict]) -> int | None:
    if not points:
        return None
    target = 10 * 60
    tolerance = 10
    scored = []
    for p in points:
        text = str(p.get("measured_at") or "")
        hhmm = text[11:16] if len(text) >= 16 else ""
        if not _valid_hhmm(hhmm):
            continue
        mins = int(hhmm[:2]) * 60 + int(hhmm[3:])
        diff = abs(mins - target)
        if diff <= tolerance:
            scored.append((diff, p))
    if not scored:
        return None
    scored.sort(key=lambda x: x[0])
    return int(scored[0][1].get("kw") or 0)


def _hourly_points(site_id: str, capacity_kw: int, date: str) -> list[dict]:
    base = max(0, int(capacity_kw * 0.62))
    seed = sum(ord(c) for c in site_id) % 700
    points: list[dict] = []
    for hour in range(10, 18):
        kw = min(capacity_kw, base + seed + (hour - 10) * 70)
        points.append({"measured_at": f"{date}T{hour:02d}:00:00+09:00", "kw": int(kw)})
    return points


def _summary_master_rows(limit_non: int = 8) -> list[dict]:
    masters = list_master_sites()
    pattern = re.compile(r"^SKK\d{3}$")
    masters = [m for m in masters if pattern.match(str(m.get("site_id") or ""))]
    by_site = {str(m.get("site_id") or ""): m for m in masters}
    canonical_semi_ids = ("SKK046", "SKK056", "SKK144", "SKK167")

    semi_rows: list[dict] = []
    for site_id in canonical_semi_ids:
        row = by_site.get(site_id)
        if row:
            semi_rows.append(row)
            continue
        semi_rows.append(
            {
                "site_id": site_id,
                "site_name": f"Semi-{site_id}",
                "group_type": "준중앙",
                "nameplate_kw": 0,
                "is_active": False,
            }
        )

    non_rows = [m for m in masters if m.get("site_id") not in set(canonical_semi_ids)]
    non_rows.sort(key=lambda x: str(x.get("site_id", "")))
    return semi_rows + non_rows[:limit_non]


def _build_summary_payload(date: str | None = None) -> dict:
    used_date = _server_date(date)
    masters = _summary_master_rows()
    store = _load_schedule_store()
    date_overrides = store["byDate"].get(used_date, {})

    sites = []
    per_site: dict[str, list[dict]] = {}
    for i, m in enumerate(masters):
        site_id = m["site_id"]
        group = "semi" if "준중앙" in m["group_type"] else "non"
        cap = int(m["nameplate_kw"])
        points = _hourly_points(site_id, cap, used_date)
        per_site[site_id] = points
        ov = date_overrides.get(site_id, {})
        site = {
            "siteId": site_id,
            "siteName": m["site_name"],
            "group": group,
            "capacityKw": cap,
            "isActive": bool(ov.get("isActive", False)),
            "isCanceled": bool(ov.get("isCanceled", False)),
            "rmccStart": str(ov.get("rmccStart", "10:00")),
            "start": str(ov.get("start", "10:20")),
            "finish": str(ov.get("finish", "10:20")),
            "minKw": int(ov.get("minKw", max(1000, int(cap * 0.25)))),
            "maxKw": int(ov.get("maxKw", max(2000, int(cap * 0.65)))) if group == "semi" else None,
            "baselineKw1000": _baseline_from_points(points) if group == "non" else None,
        }
        sites.append(site)

    active_caps = [s["capacityKw"] for s in sites if s["isActive"] and not s["isCanceled"]]
    y_max = 10000 if not active_caps else int((max(active_caps) + 9999) // 10000 * 10000)

    if sites:
        anchor = sites[0]
        summary = {
            "siteId": anchor["siteId"],
            "siteName": anchor["siteName"],
            "group": anchor["group"],
            "capacityKw": anchor["capacityKw"],
            "currentKw": per_site[anchor["siteId"]][-1]["kw"],
            "note": "live api sample",
        }
    else:
        summary = {}

    aggregate = []
    if per_site:
        for hour in range(10, 18):
            stamp = f"{used_date}T{hour:02d}:00:00+09:00"
            kw_sum = sum((pts[hour - 10]["kw"] for pts in per_site.values() if len(pts) >= (hour - 9)))
            aggregate.append({"measured_at": stamp, "kw": kw_sum})

    logs = _normalize_operation_logs(store["systemLogs"].get(used_date, []))
    return {
        "scenario": "live",
        "date": used_date,
        "timezone": "Asia/Seoul",
        "baselineTime": "10:00",
        "endTime": "17:00",
        "summary": summary,
        "sites": sites,
        "timeseriesBySiteId": per_site,
        "series": aggregate,
        "table": [
            {"label": "Date", "value": used_date},
            {"label": "Timezone", "value": "Asia/Seoul"},
            {"label": "Active sites", "value": str(len([s for s in sites if s["isActive"] and not s["isCanceled"]]))},
            {"label": "yMax", "value": f"{y_max:,} kW"},
        ],
        "systemLogs": [x["message"] for x in logs],
        "operationLogs": logs,
        "userNote": "",
        "yMaxKw": y_max,
    }


@app.get("/api/summary")
def api_summary(date: str | None = Query(default=None)) -> dict:
    return _build_summary_payload(date)


@app.patch("/api/schedule")
def api_patch_schedule(
    payload: SchedulePatchIn,
    x_user_role: str | None = Header(default=None),
    x_user_name: str | None = Header(default=None),
) -> dict:
    role = _role_guard(x_user_role, {"operator"})
    used_date = _server_date(payload.date)
    allowed = {"rmccStart", "start", "finish", "minKw", "maxKw", "isCanceled", "isActive"}
    if payload.field not in allowed:
        raise HTTPException(status_code=400, detail="unsupported field")

    value = payload.value
    if payload.field in {"rmccStart", "start", "finish"}:
        if not isinstance(value, str) or not _valid_hhmm(value):
            raise HTTPException(status_code=400, detail="invalid time format")
    elif payload.field in {"minKw", "maxKw"}:
        if isinstance(value, bool):
            raise HTTPException(status_code=400, detail="value must be number")
        try:
            value = int(float(value))
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=400, detail="value must be number") from exc
    elif payload.field in {"isCanceled", "isActive"}:
        if isinstance(value, bool):
            pass
        elif isinstance(value, str) and value.lower() in {"true", "false"}:
            value = value.lower() == "true"
        else:
            raise HTTPException(status_code=400, detail="value must be boolean")

    store = _load_schedule_store()
    date_map = store["byDate"].setdefault(used_date, {})
    site_map = date_map.setdefault(payload.siteId, {})
    old = site_map.get(payload.field)
    site_map[payload.field] = value

    old_norm = _normalize_log_value(payload.field, old)
    new_norm = _normalize_log_value(payload.field, value)
    changed = old_norm != new_norm

    log_line = None
    if changed:
        log_obj = _make_operation_log(
            date_text=used_date,
            site_id=payload.siteId,
            site_map=site_map,
            field=payload.field,
            old_value=old,
            new_value=value,
        )
        log_line = log_obj["message"]
        logs = store["systemLogs"].setdefault(used_date, [])
        logs.append(log_obj)

    _save_schedule_store(store)

    if changed:
        latest_summary = _build_summary_payload(used_date)
        latest_site = next((s for s in latest_summary.get("sites", []) if s.get("siteId") == payload.siteId), None)
        if latest_site:
            snapshot = _make_snapshot_record(used_date, latest_site, payload.field, old, value, role, x_user_name)
            _append_history(snapshot)

    return {
        "ok": True,
        "updated": {"siteId": payload.siteId, "field": payload.field, "value": value},
        "logLine": log_line,
    }


@app.get("/api/history")
def api_history(
    siteId: str = Query(...),
    from_date: str = Query(alias="from", default="2026-01-01"),
    to_date: str | None = Query(alias="to", default=None),
    limit: int = Query(default=500, ge=1, le=5000),
) -> dict:
    upper = to_date or _server_date(None)
    matched = deque(maxlen=limit)
    for rec in _iter_history_records() or []:
        rec_date = str(rec.get("date") or "")
        if rec.get("siteId") != siteId:
            continue
        if rec_date < from_date or rec_date > upper:
            continue
        matched.append(rec)
    items = list(reversed(matched))
    return {"items": items, "limit": limit}


@app.get("/api/history/summary")
def api_history_summary(
    siteId: str = Query(...),
    from_date: str = Query(alias="from", default="2026-01-01"),
    to_date: str | None = Query(alias="to", default=None),
) -> dict:
    upper = to_date or _server_date(None)
    counts = {"ORDER_CREATED": 0, "ORDER_UPDATED": 0, "ORDER_CANCELED": 0}
    for rec in _iter_history_records() or []:
        if rec.get("siteId") != siteId:
            continue
        rec_date = str(rec.get("date") or "")
        if rec_date < from_date or rec_date > upper:
            continue
        event = str(rec.get("eventType") or "")
        if event in counts:
            counts[event] += 1
    return {
        "siteId": siteId,
        "from": from_date,
        "to": upper,
        "createdCount": counts["ORDER_CREATED"],
        "updatedCount": counts["ORDER_UPDATED"],
        "canceledCount": counts["ORDER_CANCELED"],
    }


@app.get("/api/export.xlsx")
def api_export_xlsx(
    date: str | None = Query(default=None),
    x_user_role: str | None = Header(default=None),
    x_user_name: str | None = Header(default=None),
) -> Response:
    _role_guard(x_user_role, {"operator"})
    payload = _build_summary_payload(date)

    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["date", "generatedAtKst", "timezone", "yMaxKw", "operator"])
    ws.append([
        payload["date"],
        _now_kst().isoformat(timespec="seconds"),
        payload["timezone"],
        payload["yMaxKw"],
        x_user_name or "",
    ])

    semi = wb.create_sheet("SemiCentral")
    semi.append(["siteId", "siteName", "rmccStart", "start", "finish", "minKw", "maxKw", "isCanceled", "isActive"])
    for s in payload.get("sites", []):
        if s.get("group") != "semi":
            continue
        semi.append([s.get("siteId"), s.get("siteName"), _excel_kst_datetime(payload.get("date", ""), s.get("rmccStart")), _excel_kst_datetime(payload.get("date", ""), s.get("start")), _excel_kst_datetime(payload.get("date", ""), s.get("finish")), s.get("minKw"), s.get("maxKw"), s.get("isCanceled"), s.get("isActive")])

    non = wb.create_sheet("NonCentral")
    non.append(["siteId", "siteName", "rmccStart", "start", "finish", "baselineKw1000", "minKw", "isCanceled", "isActive"])
    for s in payload.get("sites", []):
        if s.get("group") != "non":
            continue
        non.append([s.get("siteId"), s.get("siteName"), _excel_kst_datetime(payload.get("date", ""), s.get("rmccStart")), _excel_kst_datetime(payload.get("date", ""), s.get("start")), _excel_kst_datetime(payload.get("date", ""), s.get("finish")), s.get("baselineKw1000") if s.get("baselineKw1000") is not None else s.get("baselineKw0900"), s.get("minKw"), s.get("isCanceled"), s.get("isActive")])

    logs_sheet = wb.create_sheet("Logs")
    logs_sheet.append(["timestampKst", "siteId", "eventType", "detail", "message"])
    for row in payload.get("operationLogs", []):
        logs_sheet.append([
            row.get("timestampKst", ""),
            row.get("siteId", ""),
            row.get("eventType", ""),
            row.get("detail", ""),
            row.get("message", ""),
        ])

    note = wb.create_sheet("UserNote")
    note.append(["note"])
    note.append([payload.get("userNote", "")])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            col_letter = col[0].column_letter
            width = max(len(str(c.value or "")) for c in col[:20]) + 2
            sheet.column_dimensions[col_letter].width = min(36, width)

    out = BytesIO()
    wb.save(out)
    filename = f"summary_{payload['date']}.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return Response(content=out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@app.post("/api/auth/login-request")
def create_login_request(payload: LoginRequestIn) -> dict:
    request_id = str(uuid4())
    now = datetime.now(tz=UTC)
    _pending_login_requests[request_id] = {
        "request_id": request_id,
        "username": payload.username,
        "device_name": payload.device_name,
        "status": "pending",
        "requested_at": now.isoformat(),
        "expires_at": (now + timedelta(minutes=3)).isoformat(),
    }
    # NOTE: password is intentionally NOT persisted in this demo flow.
    return {
        "request_id": request_id,
        "status": "pending",
        "message": "Waiting for administrator approval.",
    }


@app.get("/api/auth/pending")
def list_pending_requests() -> dict:
    now = datetime.now(tz=UTC)
    items = []
    for req in _pending_login_requests.values():
        if req["status"] == "pending" and datetime.fromisoformat(req["expires_at"]) > now:
            items.append(req)
    return {"items": items}


@app.post("/api/auth/approve/{request_id}")
def approve_login_request(request_id: str, payload: ApproveIn) -> dict:
    req = _pending_login_requests.get(request_id)
    if not req:
        raise HTTPException(status_code=404, detail="request not found")
    if req["status"] != "pending":
        raise HTTPException(status_code=409, detail="request already processed")

    now = datetime.now(tz=UTC)
    if datetime.fromisoformat(req["expires_at"]) <= now:
        req["status"] = "expired"
        raise HTTPException(status_code=410, detail="request expired")

    req["status"] = "approved"
    req["approved_by"] = payload.approver
    req["approved_at"] = now.isoformat()
    return {"request_id": request_id, "status": "approved"}


@app.post("/api/auth/token")
def exchange_token(payload: TokenExchangeIn) -> dict:
    req = _pending_login_requests.get(payload.request_id)
    if not req:
        raise HTTPException(status_code=404, detail="request not found")
    if req["status"] != "approved":
        raise HTTPException(status_code=403, detail="not approved")

    token = token_urlsafe(32)
    expires_at = datetime.now(tz=UTC) + timedelta(hours=8)
    _session_tokens[token] = {
        "username": req["username"],
        "device_name": req["device_name"],
        "expires_at": expires_at.isoformat(),
    }

    req["status"] = "completed"
    return {
        "access_token": token,
        "token_type": "bearer",
        "expires_at": expires_at.isoformat(),
    }
